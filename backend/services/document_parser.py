import io
import re
from pathlib import Path
from typing import Tuple

def parse_file(filename: str, content: bytes) -> Tuple[str, str]:
    """
    Parses uploaded file bytes and returns (extracted_text, detected_title).
    Supports PDF, DOCX, TXT, MD.
    """
    ext = Path(filename).suffix.lower()
    title = Path(filename).stem.replace("_", " ").replace("-", " ").title()
    text = ""

    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            pages = []
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    pages.append(page_text.strip())
            text = "\n\n".join(pages)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF document: {str(e)}")

    elif ext in [".docx", ".doc"]:
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            # Also extract tables if any
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n\n".join(paragraphs)
        except Exception as e:
            raise ValueError(f"Failed to parse Word (.docx) document: {str(e)}")

    elif ext in [".txt", ".md", ".markdown", ".rtf"]:
        for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            text = content.decode("utf-8", errors="ignore")

    else:
        # Check if file is an unsupported binary format
        binary_exts = [".class", ".exe", ".bin", ".dll", ".so", ".zip", ".tar", ".gz", ".pyc", ".iso", ".png", ".jpg", ".jpeg", ".mp4", ".mp3", ".jar"]
        if ext in binary_exts or b"\x00" in content[:1024]:
            raise ValueError(f"'{ext or 'binary'}' is a compiled or binary format. Please upload a readable text document (.pdf, .docx, .txt, or .md).")

        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            raise ValueError(f"Unable to decode '{ext}' document. Please upload a standard UTF-8 text file, PDF, or Word document.")

    # Remove null bytes and non-printable control characters
    clean_text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text).strip()
    if not clean_text or len(clean_text) < 10:
        raise ValueError("The uploaded document appears to be empty, binary, or unreadable.")

    # Sanitize title
    clean_title = re.sub(r'[\x00-\x1f\x7f-\x9f]', '', title).strip()
    if not clean_title or len(clean_title) < 2:
        clean_title = Path(filename).stem or "Document"

    # Infer a better title if the first line is short and heading-like
    lines = [l.strip() for l in clean_text.split("\n") if l.strip()]
    if lines and len(lines[0]) < 100 and not lines[0].endswith("."):
        candidate = lines[0].lstrip("#").strip()
        candidate = re.sub(r'[\x00-\x1f\x7f-\x9f]', '', candidate).strip()
        if candidate and len(candidate) > 3:
            clean_title = candidate

    return clean_text, clean_title
